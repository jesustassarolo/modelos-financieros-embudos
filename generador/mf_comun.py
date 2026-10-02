"""Piezas comunes a los cuatro modelos: textos de Inicio, glosario base, chequeos base, marca de agua."""

AUTOR = "Jesús Tassarolo"
MARCA_INSTAGRAM = ""   # completar con el usuario de Instagram (p. ej. "@usuario") para que aparezca en la marca de agua
MARCA_YOUTUBE = "youtube.com/@JesusTassaroloSinFiltro"
MARCA = "Modelo financiero de embudos · creado por " + AUTOR + " · TooAudience · " + MARCA_YOUTUBE + ((" · Instagram " + MARCA_INSTAGRAM) if MARCA_INSTAGRAM else "")
VERSION = "v1.2 · octubre 2026"


def inicio(nombre, que_es, ejemplo, unidad_s, unidad_p, periodo, reglas_extra=()):
    reglas = [
        f"1. Antes de invertir: cargá tus números en Supuestos (si no tenés datos, escenario 1, conservador) y mirá en Resumen el techo y el objetivo de costo por {unidad_s}.",
        f"2. Primeros días: compará tu costo real por {unidad_s} con el objetivo. Verde: seguí. Amarillo: no subas la inversión y revisá en este orden: CPM (público y frecuencia), creativo (CTR) y página (conversión). Rojo: bajá un escalón (−25 %) y arreglá antes de volver a subir.",
        f"3. Cada {periodo}: completá una columna de Seguimiento. El semáforo compara tu costo con el techo móvil (lo que rindieron tus últimas 4 columnas), no con un número fijo.",
        "4. Subí la inversión de a 25 % y esperá 7 días (o 3 a 4 eventos si vendés por llamada) antes de volver a subir. Un día bueno o malo no decide nada.",
        "5. Si el costo está en verde y aun así no hay ganancia, el problema no es el tráfico: mirá asistencia, conversión de la oferta y cobranza, en ese orden.",
        "6. Cambiá una sola cosa por vez. Si cambiás precio, creativo y página la misma semana no vas a saber qué funcionó.",
        "7. Trabajá siempre con cash neto (lo que entra menos comisiones y reembolsos), nunca con facturado. En Proyección ves cuánta caja necesitás para sostener la pauta mientras cobrás.",
        *reglas_extra,
    ]
    return [
        ("Para qué sirve este modelo", [
            f"Es el modelo financiero de un {nombre}. Con tus números (o con el ejemplo cargado) te dice cuánto vale cada {unidad_s} que comprás con publicidad, hasta cuánto podés pagar por {unidad_s} y por cada etapa siguiente, cuánto invertir para una meta de ventas, cuánta caja necesitás y qué palanca mejora más la ganancia.",
            *que_es,
        ]),
        ("Cómo usarlo en 7 pasos", [
            "1. Andá a Supuestos y reemplazá el ejemplo por tus números. Si todavía no tenés datos, dejá el escenario en 1 (conservador) y usá Benchmarks como referencia.",
            f"2. Mirá Resumen: ganancia del {periodo}, ROAS, el techo y el objetivo de costo por {unidad_s}, y el semáforo. Después Embudo: todos los pasos en orden con cuántos pasan, lo que pagás, lo que te queda y lo máximo que podrías pagar por cada uno.",
            "3. Para jugar con los números, usá el Excel de Simulación de este embudo (archivo aparte): es esta misma cadena, idéntica a la hoja Modelo, con todos los supuestos editables en su lugar. Cambiá uno y se recalcula todo. No se conecta con este archivo: cuando una prueba te convence, pasá ese número a Supuestos.",
            "4. Abrí Escenarios para ver qué pasa si cambia el precio, la conversión o el costo en tablas de dos variables, y Sensibilidad para saber qué variable conviene trabajar primero.",
            "5. Antes de invertir, mirá Proyección 12 meses: la caja necesaria es la plata que tenés que tener para no frenar la pauta mientras cobrás. Si da 0, el embudo se autofinancia.",
            f"6. Cuando el embudo esté andando, completá una columna de Seguimiento por {periodo}. El semáforo te dice si estás pagando de más por cada {unidad_s}.",
            "7. Una decisión por semana, anotada. A los 7 días, mirá qué pasó.",
        ]),
        ("Colores y hojas", [
            "Amarillo = celda que editás vos. Azul = resultado calculado: no lo toques, se recalcula solo.",
            "Resumen, Embudo y Modelo leen de Supuestos. Escenarios y Sensibilidad recalculan el modelo entero cambiando una o dos variables. Proyección usa los mismos supuestos mes a mes. Seguimiento es la única hoja, además de Supuestos, donde cargás datos.",
            "Chequeos avisa si un supuesto quedó fuera de rango. Si dice REVISAR, arreglalo antes de leer los resultados.",
            "Cómo elegir escenario: en Supuestos, la celda 'Escenario en uso' (1, 2 o 3) elige la columna conservador, base u optimista de cada fila que tiene tres valores. Las filas con un solo valor se usan tal cual.",
        ]),
        ("El ejemplo que viene cargado", ejemplo),
        ("Reglas de decisión (el protocolo corto)", reglas),
        ("Lo que este modelo no hace", [
            "No adivina tus conversiones: usa las que cargás. Cuando una conversión cambia, el techo cambia; por eso Seguimiento recalcula el techo móvil con tus últimas 4 columnas.",
            "No reemplaza la decisión: la informa. Las reglas de arriba dicen qué hacer con cada color.",
            "Los benchmarks son orden de magnitud para tráfico frío en español. Tus datos mandan.",
        ]),
        ("Autor y versión", [MARCA + ". " + VERSION + ". Fórmulas compatibles con Excel, Google Sheets y Numbers. Sin macros.",
                            "Este modelo se comparte para que lo uses con tus números. Si lo compartís, dejá esta línea: es la forma de saber de dónde salió."]),
    ]


GLOSARIO_BASE = [
    ("CPM", "Costo por mil impresiones: lo que cobra la plataforma de anuncios por mostrar el aviso mil veces.", "Supuestos"),
    ("CTR", "Porcentaje de impresiones que terminan en clic al enlace.", "Supuestos"),
    ("Clic → visita", "Porcentaje de clics que llegan a cargar la página. Siempre menor a 100 % (cierran antes de que cargue).", "Supuestos"),
    ("Unidad de adquisición", "Lo que compra la publicidad en este embudo (visita, registro o entrada). Todo el modelo se mide por unidad.", "Modelo"),
    ("Costo por unidad", "Inversión dividida por unidades compradas: costo por visita, CPL (costo por lead o registro) o costo por entrada.", "Modelo, Seguimiento"),
    ("Conversión", "Porcentaje que pasa de una etapa a la siguiente. Las conversiones se multiplican: 15 % × 14 % × 44 % = 0,9 %.", "Supuestos, Modelo"),
    ("Show rate", "Porcentaje de registrados (o compradores) que asisten en vivo. El replay se cuenta aparte.", "Supuestos"),
    ("Bump", "Oferta extra en el mismo checkout, con una casilla. Se mide como % de los compradores.", "Supuestos"),
    ("OTO / upsell", "Oferta que aparece después de pagar (one-time offer). Downsell: la oferta más barata para quien rechaza la OTO.", "Supuestos"),
    ("Backend", "Producto de mayor valor que se vende después del primero, por ascenso o por llamada.", "Supuestos"),
    ("AOV", "Facturado promedio por comprador del producto principal, sumando toda la escalera.", "Modelo"),
    ("CPA / CAC", "Costo por comprador: inversión dividida por ventas del producto principal.", "Modelo"),
    ("Facturado", "Lo vendido a precio de lista, aunque todavía no se haya cobrado (cuotas).", "Modelo"),
    ("Cash cobrado", "Lo que efectivamente entró. Con cuotas, entra en varios meses según la curva de cobranza.", "Modelo, Proyección"),
    ("Cash neto", "Cash cobrado menos comisiones de pasarela, comisiones de ventas y reembolsos. Es la base de verdad del modelo.", "Modelo"),
    ("ROAS", "Retorno sobre la inversión en publicidad. Sobre facturado, sobre cash o sobre cash neto: siempre decí cuál.", "Modelo"),
    ("Valor neto por unidad (VPU)", "Cash neto que deja cada unidad de adquisición, descontando los costos que crecen con el volumen. Es lo que vale un lead, una visita o una entrada.", "Modelo"),
    ("Techo", "Costo por unidad al que no ganás ni perdés (break-even). Es igual al valor neto por unidad.", "Modelo, Seguimiento"),
    ("Objetivo", "Techo dividido por (1 + margen de seguridad). Si pagás eso o menos, estás en verde.", "Modelo"),
    ("Margen de seguridad", "Colchón entre el techo y el objetivo para absorber varianza, reembolsos y costos que no viste. 30 % es el estándar.", "Supuestos"),
    ("Semáforo", "Rojo: el costo por unidad está por encima del techo (perdés con cada unidad). Verde: el costo está en el objetivo o por debajo y el ROAS sobre cash neto llega al objetivo. Amarillo: todo lo demás (el costo está entre objetivo y techo, o el ROAS no llega).", "Resumen, Seguimiento"),
    ("Techo móvil", "Mediana del valor neto por unidad de tus últimas 4 columnas de Seguimiento. Es el techo que se usa para decidir, porque refleja tu rendimiento reciente.", "Seguimiento"),
    ("Escenario", "Conjunto de supuestos conservador, base u optimista. Se elige con una celda en Supuestos.", "Supuestos"),
    ("Sensibilidad", "Cuánto cambia la ganancia si mejorás una sola variable un 10 %. Dice qué palanca trabajar primero.", "Sensibilidad"),
    ("Ganancia por visita / por lead / por cliente", "La ganancia del período dividida por la cantidad de ese paso. Es la otra cara del costo por paso. La ganancia por visita es la métrica que más vale: resume todo lo que pasa después de la página.", "Modelo, Embudo"),
    ("Order value (AOV)", "Facturado promedio por comprador del principal sumando toda la escalera. El AOV neto es lo mismo después de comisiones y reembolsos, y su par es el CPA.", "Modelo"),
    ("Caja necesaria", "El peor momento de la ganancia acumulada en la proyección, en positivo. Es la plata que necesitás tener para no frenar la pauta mientras cobrás. Si da 0, el embudo se autofinancia.", "Proyección"),
    ("Unit economics", "Los números por unidad: cuánto cuesta y cuánto deja cada visita, registro, entrada o comprador.", "Resumen"),
    ("Tráfico frío", "Gente que no te conoce y llega desde un anuncio. Convierte menos que tu audiencia propia.", "Benchmarks"),
    ("Front / producto principal", "El primer producto que se vende en el embudo. Todo lo demás (bumps, OTOs, programa) se mide como % de sus compradores.", "Supuestos"),
    ("Pitch", "La parte del webinar o la llamada donde se presenta la oferta.", "Inicio"),
    ("VSL", "Video de ventas: la página de ventas es un video largo que vende el producto.", "Supuestos"),
    ("B2B", "Negocios que venden a empresas. Sus benchmarks de asistencia son más altos que los de tráfico frío a personas.", "Benchmarks"),
    ("Self-serve", "Venta sin llamada: la persona compra sola desde una página de pago.", "Supuestos"),
]


def chequeos_base(tasas, cadena, positivos=(), extra=()):
    """tasas: claves entre 0 y 1. cadena: pares (mayor, menor) con menor <= mayor. positivos: claves que deben ser > 0. extra: chequeos propios."""
    chs = [
        {"label": "Los supuestos de tráfico son mayores a cero", "cond": "AND(" + ",".join(f"{{{k}}}>0" for k in positivos) + ")", "ayuda": "Con un 0 en el tráfico o en la conversión de la página no hay unidades y el modelo no dice nada."} if positivos else None,
        {"label": "La curva de cobranza suma 100 %", "cond": "ABS({cobro_m1}+{cobro_m2}+{cobro_m3}-1)<0.001", "ayuda": "Supuestos → cobranza mes 1 + mes 2 + mes 3 = 100 %."},
        {"label": "Comisiones + reembolsos menores al 60 %", "cond": "({pasarela}+{closers}+{reembolsos})<0.6", "ayuda": "Si suman más, revisá los porcentajes."},
        {"label": "El escenario en uso es 1, 2 o 3", "cond": "OR({escenario}=1,{escenario}=2,{escenario}=3)", "ayuda": "Supuestos → celda 'Escenario en uso'."},
        {"label": "La inversión es mayor a cero", "cond": "{inversion}>0", "ayuda": ""},
        {"label": "La meta de ventas es alcanzable (hay conversión a venta)", "cond": "{conv_unidad_venta}>0", "ayuda": "Con conversión a venta en 0, la inversión para la meta dice 'inalcanzable con estos supuestos'."},
        {"label": "Las conversiones están entre 0 y 100 %", "cond": "AND(" + ",".join(f"{{{k}}}>=0,{{{k}}}<=1" for k in tasas) + ")", "ayuda": "Un 25 % se escribe 25 % (o 0,25), no 25."},
        {"label": "El margen de seguridad está entre 0 y 100 %", "cond": "AND({margen_seg}>=0,{margen_seg}<=1)", "ayuda": ""},
        {"label": "El crecimiento mensual está entre −50 % y 100 %", "cond": "AND({crecimiento}>=-0.5,{crecimiento}<=1)", "ayuda": "Más de 100 % por mes no es realista para una proyección."},
        {"label": "La cadena es decreciente (cada etapa tiene menos gente que la anterior)", "cond": "AND(" + ",".join(f"{{{b}}}<={{{a}}}" for a, b in cadena) + ")", "ayuda": "Si falla, alguna conversión está por encima de 100 %."},
        *extra,
    ]
    return [c for c in chs if c]


def cobranza(m1, m2, m3, nota):
    return [
        {"key": "cobro_m1", "label": "% del facturado que entra en el mes 1", "val": m1, "fmt": "pct", "unidad": "%", "nota": nota},
        {"key": "cobro_m2", "label": "% que entra en el mes 2", "val": m2, "fmt": "pct", "unidad": "%", "nota": "Cuotas o pagos diferidos."},
        {"key": "cobro_m3", "label": "% que entra en el mes 3", "val": m3, "fmt": "pct", "unidad": "%", "nota": "Los tres deben sumar 100 %."},
    ]


def objetivos(roas, meta, meta_unidad, caja, periodo):
    return [
        {"key": "roas_obj", "label": "ROAS objetivo sobre cash neto", "val": roas, "fmt": "ratio", "unidad": "x", "nota": "El retorno mínimo que querés por cada dólar invertido, después de comisiones y reembolsos. Entra en el semáforo: sin llegar a este ROAS no hay verde."},
        {"key": "margen_seg", "label": "Margen de seguridad", "val": 0.30, "fmt": "pct", "unidad": "%", "nota": "Colchón entre el techo (donde no ganás ni perdés) y el objetivo. 30 % es el estándar que usamos."},
        {"key": "meta_ventas", "label": "Meta de ventas del producto principal", "val": meta, "fmt": "int", "unidad": meta_unidad, "nota": "Para calcular cuánto invertir."},
        {"key": "crecimiento", "label": "Crecimiento de la inversión por mes (proyección)", "val": 0.10, "fmt": "pct", "unidad": "% / mes", "nota": "Cuánto sube la inversión cada mes en la proyección. 10 % por mes es un ritmo prudente; más de 25 % suele encarecer el CPM."},
        {"key": "caja_inicial", "label": "Caja inicial", "val": caja, "fmt": "money0", "unidad": "USD", "nota": "Con cuánta plata arrancás. La proyección te dice si alcanza."},
    ]


def costos(pasarela, closers, reemb, semivar, semivar_unidad, fijos, periodo, nota_closers, nota_semivar):
    return [
        {"key": "pasarela", "label": "Comisión de la pasarela o plataforma de pago", "val": pasarela, "fmt": "pct", "unidad": "% del cash", "nota": "Hotmart, Stripe, Mercado Pago: lo que se queda la plataforma."},
        {"key": "closers", "label": "Comisión del equipo de ventas", "val": closers, "fmt": "pct", "unidad": "% del cash", "nota": nota_closers},
        {"key": "reembolsos", "label": "Reembolsos y contracargos", "val": reemb, "fmt": "pct", "unidad": "% del cash", "nota": "5 % es normal en productos digitales. Más de 10 % es un problema de oferta o de expectativa."},
        {"key": "costo_semivar", "label": f"Costos por {semivar_unidad} (herramientas, WhatsApp, IA)", "val": semivar, "fmt": "money", "unidad": f"USD / {semivar_unidad}", "nota": nota_semivar},
        {"key": "fijos", "label": f"Costos fijos por {periodo}", "val": fijos, "fmt": "money0", "unidad": f"USD / {periodo}", "nota": "Equipo, herramientas, edición: lo que pagás aunque no vendas."},
    ]


def cash_rows(unidad_label):
    """Filas de cash, costos y ganancia, iguales en los cuatro modelos."""
    return [
        {"seccion": "Cash, costos y ganancia del período"},
        {"key": "cash", "label": "Cash a cobrar (facturado × cobranza completa)", "f": "{facturado}*({cobro_m1}+{cobro_m2}+{cobro_m3})", "fmt": "money0", "como": "Facturado × (mes 1 + mes 2 + mes 3). Si la cobranza suma 100 %, es igual al facturado."},
        {"key": "cash_m1", "label": "Cash que entra en el primer mes", "f": "{facturado}*{cobro_m1}", "fmt": "money0", "como": "Facturado × % del mes 1."},
        {"key": "comisiones", "label": "Comisiones (pasarela + ventas)", "f": "{cash}*({pasarela}+{closers})", "fmt": "money0", "como": "Cash × (pasarela + equipo de ventas)."},
        {"key": "reembolsos_usd", "label": "Reembolsos", "f": "{cash}*{reembolsos}", "fmt": "money0", "como": "Cash × % de reembolsos."},
        {"key": "neto", "label": "Cash neto", "f": "{cash}*(1-{pasarela}-{closers}-{reembolsos})", "fmt": "money0", "como": "Cash − comisiones − reembolsos. La base de verdad.", "bold": True},
        {"key": "semivar", "label": f"Costos por {unidad_label} (herramientas, WhatsApp, IA)", "f": "{unidades}*{costo_semivar}", "fmt": "money0", "como": f"{unidad_label.capitalize()}s × costo por {unidad_label}."},
        {"key": "profit", "label": "Ganancia del período", "f": "{neto}-{semivar}-{fijos}-{inversion}", "fmt": "money0", "como": "Cash neto − costos por unidad − fijos − inversión.", "bold": True},
        {"key": "roas_fact", "label": "ROAS sobre facturado", "f": "{facturado}/{inversion}", "fmt": "ratio", "como": "Facturado ÷ inversión. El número que más se infla: no decidas con este."},
        {"key": "roas_cash", "label": "ROAS sobre cash neto", "f": "{neto}/{inversion}", "fmt": "ratio", "como": "Cash neto ÷ inversión. Con este se decide.", "bold": True},
        {"key": "margen", "label": "Margen sobre cash", "f": "IF({cash}=0,0,{profit}/{cash})", "fmt": "pct", "como": "Ganancia ÷ cash."},
        {"key": "cumple_roas", "label": "¿Cumple el ROAS objetivo?", "f": 'IF({roas_cash}>={roas_obj},"SÍ","NO")', "fmt": "txt", "como": "ROAS sobre cash neto contra el objetivo de Supuestos."},
    ]


def seguimiento_calc(inv_key, unidades_key, fact_key, cash_key):
    """Filas calculadas comunes de Seguimiento. El valor por unidad y el techo se calculan sobre el facturado
    (cobranza completa, como en Modelo); la ganancia de caja, sobre lo cobrado hasta ahora."""
    neto_f = f"{{{fact_key}}}*(1-{{pasarela}}-{{closers}}-{{reembolsos}})"
    neto_c = f"{{{cash_key}}}*(1-{{pasarela}}-{{closers}}-{{reembolsos}})"
    u, inv = f"{{{unidades_key}}}", f"{{{inv_key}}}"
    return [
        {"key": "vpu_r", "label": "Valor neto por unidad real (sobre facturado, cobranza completa)", "tipo": "f", "f": f"IF({u}=0,0,({neto_f}-{u}*{{costo_semivar}})/{u})", "fmt": "money", "bold": True},
        {"key": "profit_r", "label": "Ganancia del período al completar la cobranza", "tipo": "f", "f": f"{neto_f}-{u}*{{costo_semivar}}-{{fijos}}-{inv}", "fmt": "money0", "bold": True},
        {"key": "profit_caja_r", "label": "Ganancia de caja del período (con lo cobrado hasta ahora)", "tipo": "f", "f": f"{neto_c}-{u}*{{costo_semivar}}-{{fijos}}-{inv}", "fmt": "money0"},
        {"key": "gan_unidad_r", "label": "Ganancia por unidad real (ganancia ÷ unidades)", "tipo": "f", "f": f"IF({u}=0,0,({neto_f}-{u}*{{costo_semivar}}-{{fijos}}-{inv})/{u})", "fmt": "money", "bold": True},
        {"key": "roas_r", "label": "ROAS sobre cash neto (cobranza completa)", "tipo": "f", "f": f"IF({inv}=0,0,{neto_f}/{inv})", "fmt": "ratio", "bold": True},
        {"key": "roas_caja_r", "label": "ROAS de caja (cobrado hasta ahora)", "tipo": "f", "f": f"IF({inv}=0,0,{neto_c}/{inv})", "fmt": "ratio"},
        {"key": "techo_r", "label": "Techo móvil (mediana de tus últimas 4 columnas; al inicio, el del modelo)", "tipo": "techo_movil", "src": "vpu_r", "fmt": "money"},
        {"key": "sem_r", "label": "Semáforo (costo por unidad contra el techo móvil, y ROAS contra el objetivo)", "tipo": "semaforo", "cpu": "cpu_r", "techo": "techo_r", "roas": "roas_r", "fmt": "txt", "bold": True},
    ]


# ------------------------------------------------------------ marca de agua: metadatos, encabezados de impresión, fila 1 y hoja Licencia
LICENCIA = ("Licencia Creative Commons BY-NC-ND 4.0: podés usar este archivo con tus números y compartirlo tal cual, citando al autor. "
            "No podés venderlo, modificar su estructura para redistribuirlo ni quitar la línea de autoría.")


def aplicar_marca(wb, titulo, descripcion_extra=""):
    """Deja la autoría de Jesús Tassarolo en todo el libro: propiedades (autor, título, asunto, palabras clave, descripción),
    propiedades personalizadas, nombre definido, encabezado y pie de impresión de cada hoja, una línea en la fila 1 de cada
    hoja y una hoja Licencia al final. Google Sheets conserva la fila 1, la hoja Licencia y el título; Excel y Numbers,
    además, los metadatos y los encabezados de impresión."""
    from openpyxl.packaging.custom import StringProperty
    from openpyxl.workbook.defined_name import DefinedName
    from openpyxl.styles import Font
    p = wb.properties
    p.creator = AUTOR; p.lastModifiedBy = AUTOR; p.title = titulo
    p.subject = "Modelo financiero de embudos · " + AUTOR + " · TooAudience"
    p.description = MARCA + " · " + LICENCIA + (" · " + descripcion_extra if descripcion_extra else "")
    p.keywords = AUTOR + ", TooAudience, modelo financiero, embudos, mastermind" + ((", Instagram " + MARCA_INSTAGRAM) if MARCA_INSTAGRAM else "")
    p.category = "TooAudience · Mastermind"
    for nombre, valor in [("Autor", AUTOR), ("Marca", MARCA), ("Licencia", "CC BY-NC-ND 4.0"), ("YouTube", MARCA_YOUTUBE), ("Version", VERSION)] + ([("Instagram", MARCA_INSTAGRAM)] if MARCA_INSTAGRAM else []):
        try:
            wb.custom_doc_props.append(StringProperty(name=nombre, value=valor))
        except Exception:
            pass
    try:
        wb.defined_names["Autor_Jesus_Tassarolo"] = DefinedName("Autor_Jesus_Tassarolo", attr_text='"' + MARCA.replace('"', "'") + '"')
    except Exception:
        pass
    nota = Font(italic=True, color="8C8C8C", size=8)
    for ws in wb.worksheets:
        if ws.title == "Licencia":
            continue
        ws.oddHeader.center.text = AUTOR + " · TooAudience"
        ws.oddHeader.right.text = "&D"
        ws.oddFooter.left.text = MARCA[:250]
        ws.oddFooter.right.text = "Página &P de &N"
        ws.evenHeader.center.text = AUTOR + " · TooAudience"; ws.evenFooter.left.text = MARCA[:250]; ws.evenFooter.right.text = "Página &P de &N"
        c = 2 if (ws.cell(2, 2).value is not None and ws.cell(2, 1).value is None) else 1
        if ws.cell(1, c).value is None:
            ws.cell(1, c, "© " + AUTOR + " · TooAudience · " + MARCA_YOUTUBE + " · uso personal; no redistribuir modificado").font = nota
    if "Licencia" not in wb.sheetnames:
        ws = wb.create_sheet("Licencia"); ws.column_dimensions["A"].width = 3; ws.column_dimensions["B"].width = 120
        ws.cell(2, 2, "Autoría y licencia").font = Font(bold=True, size=14)
        for i, t in enumerate([titulo, MARCA, VERSION, "", LICENCIA, "",
                               "Este archivo, sus fórmulas, sus umbrales y su método son obra de " + AUTOR + " (TooAudience). Si lo compartís, compartilo entero y con esta hoja.",
                               "Canal: " + MARCA_YOUTUBE + ((" · Instagram " + MARCA_INSTAGRAM) if MARCA_INSTAGRAM else "")]):
            ws.cell(4 + i, 2, t).alignment = __import__("openpyxl").styles.Alignment(wrap_text=True, vertical="top")
        ws.sheet_properties.tabColor = "1F3864"
    return wb
