#!/usr/bin/env python3
"""Verifica los cuatro libros: recalcula todas las fórmulas con el motor `formulas` (sin Excel),
busca errores (#DIV/0!, #NAME?, #REF!...), compara los resultados clave contra cuentas hechas a mano
(independientes de las especificaciones) y chequea la coherencia entre hojas.

Uso:  python3 verificar.py [carpeta_entregables]
"""
import os, re, sys, glob
import openpyxl
import formulas

HERE = os.path.dirname(os.path.abspath(__file__))
DEFAULT_IN = os.path.normpath(os.path.join(HERE, "..", "excel"))
CELL_RE = re.compile(r"^'?\[.*?\](.+?)'?!\$?([A-Z]+)\$?(\d+)$")

# Cuentas a mano del escenario base de cada ejemplo (ver Inicio de cada libro).
def esperado_low_ticket():
    inv, cpu = 9000, 0.30
    vis = inv / cpu; ck = vis * 0.035; v = ck * 0.26
    fact = v * (37 + 0.25 * 17 + 0.18 * 17 + 0.10 * 47 + 0.06 * 17 + 0.03 * 97 + 0.04 * 27 + 0.02 * 297)
    neto = fact * (1 - 0.10 - 0.05); profit = neto - 1500 - inv
    return {"Visitas a la página de ventas": vis, "Ventas del producto principal": v, "Facturado total": fact,
            "Cash neto": neto, "Ganancia del período": profit, "Valor neto por visita": neto / vis,
            "ROAS sobre cash neto": neto / inv, "Costo por comprador (CPA)": inv / v}

def esperado_webinar():
    inv = 6000; cpl = 4.5 / 1000 / (0.022 * 0.9 * 0.23); leads = inv / cpl
    vivo, rep = leads * 0.15, leads * 0.55
    ventas = (vivo * 0.10 + rep * 0.025) * 0.40
    fact = ventas * (297 + 0.20 * 47 + 0.05 * 997)
    neto = fact * 0.85; semi = leads * 0.10; profit = neto - semi - 1200 - inv
    return {"Costo por registro (CPL)": cpl, "Registros (leads)": leads, "Asistentes en vivo": vivo,
            "Ventas de la oferta": ventas, "Facturado total": fact, "Cash neto": neto,
            "Ganancia del período": profit, "Valor neto por registro": (neto - semi) / leads, "ROAS sobre cash neto": neto / inv}

def esperado_llamada():
    inv = 4500; cpl = 10 / 1000 / (0.015 * 0.85 * 0.28); leads = inv / cpl
    apl = leads * 0.10; cal = apl * 0.60; ag = cal * 0.60; ll = ag * 0.70; v = ll * 0.25; ds = (ll - v) * 0.10
    fact = v * 1500 + ds * 497; neto = fact * (1 - 0.05 - 0.10 - 0.05); semi = leads * 0.20
    profit = neto - semi - 2500 - inv
    return {"Costo por lead (CPL)": cpl, "Leads": leads, "Agendas": ag, "Llamadas realizadas": ll,
            "Ventas del programa": v, "Facturado total": fact, "Cash neto": neto, "Ganancia del período": profit,
            "Valor neto por lead": (neto - semi) / leads, "ROAS del cash del primer mes (sin cuotas futuras)": fact * 0.6 * 0.8 / inv}

def esperado_pago():
    inv = 3000; vis = inv / 0.35; ent = vis * 0.025
    front = ent * 9 + ent * 0.20 * 27
    vivo, rep = ent * 0.55, ent * 0.25; v = (vivo * 0.16 + rep * 0.04) * 0.50
    fact = front + v * 497 + v * 0.10 * 1997; neto = fact * (1 - 0.08 - 0.05); semi = ent * 0.30
    profit = neto - semi - 800 - inv
    return {"Entradas vendidas": ent, "Costo por entrada vendida": 0.35 / 0.025, "ROAS del front (qué parte de la pauta recupera la entrada)": front / inv,
            "Ventas de la oferta": v, "Facturado total": fact, "Cash neto": neto, "Ganancia del período": profit,
            "Valor neto por entrada": (neto - semi) / ent, "ROAS sobre cash neto": neto / inv}

ESPERADOS = {"01_": esperado_low_ticket, "02_": esperado_webinar, "03_": esperado_llamada, "04_": esperado_pago}


def calcular(path):
    xl = formulas.ExcelModel().loads(path).finish()
    sol = xl.calculate()
    vals = {}
    for k, v in sol.items():
        m = CELL_RE.match(k)
        if not m:
            continue
        sheet, c, r = m.group(1).upper().strip("'"), m.group(2), int(m.group(3))
        try:
            val = v.value[0][0]
        except Exception:
            val = v
        vals[(sheet, f"{c}{r}")] = val
    return vals


def es_error(v):
    s = str(v)
    return s.startswith("#") or "XlError" in type(v).__name__


def verificar(path):
    nombre = os.path.basename(path)
    print(f"\n=== {nombre}")
    wb = openpyxl.load_workbook(path)  # fórmulas y etiquetas
    vals = calcular(path)
    fallas = []
    # 1) errores en cualquier celda con fórmula
    n_f = 0
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    n_f += 1
                    v = vals.get((ws.title.upper(), c.coordinate))
                    if v is None:
                        fallas.append(f"sin valor calculado: {ws.title}!{c.coordinate} {c.value[:60]}")
                    elif es_error(v):
                        fallas.append(f"error {v}: {ws.title}!{c.coordinate} {c.value[:80]}")
    print(f"  fórmulas evaluadas: {n_f}")
    # 2) resultados clave vs cuentas a mano
    modelo = wb["Modelo"]
    por_label = {}
    for r in range(1, modelo.max_row + 1):
        lab = modelo.cell(r, 2).value
        if lab:
            por_label[lab] = vals.get(("MODELO", f"C{r}"))
    esp = next(fn for pref, fn in ESPERADOS.items() if nombre.startswith(pref))()
    for lab, e in esp.items():
        got = por_label.get(lab)
        if got is None:
            fallas.append(f"falta la fila '{lab}' en Modelo"); continue
        ok = abs(float(got) - e) <= 0.01 * max(abs(e), 1e-9) + 1e-6
        print(f"  {'OK ' if ok else 'XX '} {lab:<58} hoja {float(got):>12,.2f}   a mano {e:>12,.2f}")
        if not ok:
            fallas.append(f"desvío en '{lab}': hoja {got} vs a mano {e}")
    # 3) coherencia entre hojas
    profit = float(por_label["Ganancia del período"])
    sens = wb["Sensibilidad"]
    base_row = next(r for r in range(1, sens.max_row + 1) if str(sens.cell(r, 2).value).startswith("Ganancia del período con los supuestos"))
    base = vals.get(("SENSIBILIDAD", f"D{base_row}"))
    if abs(float(base) - profit) > 0.01:
        fallas.append(f"Sensibilidad base {base} ≠ Modelo {profit}")
    deltas = [vals.get(("SENSIBILIDAD", f"E{r}")) for r in range(base_row + 1, sens.max_row + 1) if isinstance(sens.cell(r, 5).value, str)]
    ranks = [vals.get(("SENSIBILIDAD", f"G{r}")) for r in range(base_row + 1, sens.max_row + 1) if isinstance(sens.cell(r, 7).value, str)]
    print(f"  sensibilidad: diferencias {[round(float(d)) for d in deltas]} ranking {[int(float(x)) for x in ranks]}")
    # las tablas de Escenarios: la celda cuyo encabezado de fila y columna coinciden con el valor en uso debe dar la ganancia del modelo
    esc = wb["Escenarios"]
    coincidencias = 0
    for r in range(1, esc.max_row + 1):
        for c in range(3, esc.max_column + 1):
            cell = esc.cell(r, c)
            if isinstance(cell.value, str) and cell.value.startswith("=") and "Supuestos" in cell.value and "$B" in cell.value:
                v = vals.get(("ESCENARIOS", cell.coordinate))
                if v is not None and not es_error(v) and abs(float(v) - profit) < 0.01:
                    coincidencias += 1
    print(f"  celdas de Escenarios que reproducen la ganancia del modelo: {coincidencias}")
    # 4) chequeos
    chq = wb["Chequeos"]
    estados = [vals.get(("CHEQUEOS", f"C{r}")) for r in range(1, chq.max_row + 1) if isinstance(chq.cell(r, 3).value, str) and chq.cell(r, 3).value.startswith("=")]
    malos = [e for e in estados if str(e) != "OK"]
    print(f"  chequeos: {len(estados)} evaluados, {len(malos)} en REVISAR")
    if malos:
        fallas.append(f"chequeos en REVISAR: {len(malos)}")
    # 5) seguimiento: semáforos de las 3 columnas de ejemplo
    seg = wb["Seguimiento"]
    sem = [vals.get(("SEGUIMIENTO", f"{col}{r}")) for r in range(1, seg.max_row + 1) if seg.cell(r, 2).value and "Semáforo" in str(seg.cell(r, 2).value) for col in "CDE"]
    print(f"  semáforos de Seguimiento (ejemplo): {sem}")
    if any(str(s) not in ("VERDE", "AMARILLO", "ROJO", "") for s in sem):
        fallas.append(f"semáforo de Seguimiento inválido: {sem}")
    # 6) proyección
    pro = wb["Proyección 12 meses"]
    caja = [(str(pro.cell(r, 2).value)[:28], round(float(vals.get(("PROYECCIÓN 12 MESES", f"C{r}")) or 0))) for r in range(1, pro.max_row + 1) if pro.cell(r, 2).value and ("Caja necesaria" in str(pro.cell(r, 2).value) or "Saldo mínimo" in str(pro.cell(r, 2).value))]
    print(f"  caja: {caja}")
    # 7) estado general de Chequeos y caso borde: un supuesto de tráfico en 0 tiene que dar REVISAR
    est_row = next(r for r in range(1, chq.max_row + 1) if str(chq.cell(r, 2).value).startswith("Estado general"))
    estado = vals.get(("CHEQUEOS", f"C{est_row}"))
    print(f"  estado general de Chequeos: {estado}")
    if str(estado) != "OK":
        fallas.append("estado general de Chequeos no es OK")
    import tempfile
    tmp = os.path.join(tempfile.gettempdir(), "borde_" + nombre)
    wb2 = openpyxl.load_workbook(path); sup = wb2["Supuestos"]
    for r in range(1, sup.max_row + 1):
        lab = str(sup.cell(r, 2).value)
        if lab.startswith("CTR") or lab.startswith("Costo por visita"):
            if isinstance(sup.cell(r, 3).value, str) and str(sup.cell(r, 3).value).startswith("="):
                sup.cell(r, 5).value = 0   # columna "Base" del escenario
            else:
                sup.cell(r, 3).value = 0
            break
    wb2.save(tmp)
    v2 = calcular(tmp)
    est2 = v2.get(("CHEQUEOS", f"C{est_row}"))
    n_err = sum(1 for k, v in v2.items() if k[0] in ("MODELO", "RESUMEN") and es_error(v))
    fallas += verificar_simulador(path, vals, wb)
    print(f"  caso borde (tráfico en 0): estado general {est2}, celdas con error en Modelo/Resumen {n_err}")
    if str(est2) != "REVISAR":
        fallas.append(f"caso borde: Chequeos dice {est2} con un supuesto de tráfico en 0")
    # 8) los tres escenarios recalculan sin errores y con Chequeos en OK
    sel_row = next(r for r in range(1, sup.max_row + 1) if str(sup.cell(r, 2).value).startswith("Escenario en uso"))
    for esc in (1, 2, 3):
        wb3 = openpyxl.load_workbook(path); wb3["Supuestos"].cell(sel_row, 3).value = esc
        t3 = os.path.join(tempfile.gettempdir(), f"esc{esc}_" + nombre); wb3.save(t3)
        v3 = calcular(t3)
        e3 = v3.get(("CHEQUEOS", f"C{est_row}"))
        n3 = sum(1 for k, v in v3.items() if k[0] in ("MODELO", "RESUMEN", "PROYECCIÓN 12 MESES", "ESCENARIOS") and es_error(v))
        prof3 = v3.get(("MODELO", next(f"C{r}" for r in range(1, modelo.max_row + 1) if modelo.cell(r, 2).value == "Ganancia del período")))
        print(f"  escenario {esc}: Chequeos {e3}, errores {n3}, ganancia {float(prof3):,.0f}")
        if str(e3) != "OK" or n3:
            fallas.append(f"escenario {esc}: Chequeos {e3}, {n3} errores")
    # 9) conversión a venta en 0: la meta debe decir 'inalcanzable' y Chequeos REVISAR
    wb4 = openpyxl.load_workbook(path); sup4 = wb4["Supuestos"]
    for r in range(1, sup4.max_row + 1):
        lab = str(sup4.cell(r, 2).value)
        if "que compran" in lab or "terminan en compra" in lab:
            if isinstance(sup4.cell(r, 3).value, str) and str(sup4.cell(r, 3).value).startswith("="):
                for cc in (4, 5, 6):
                    sup4.cell(r, cc).value = 0
            else:
                sup4.cell(r, 3).value = 0
    t4 = os.path.join(tempfile.gettempdir(), "meta0_" + nombre); wb4.save(t4)
    v4 = calcular(t4)
    meta_row = next(r for r in range(1, modelo.max_row + 1) if str(modelo.cell(r, 2).value).startswith("Inversión necesaria para la meta"))
    meta_val = v4.get(("MODELO", f"C{meta_row}")); e4 = v4.get(("CHEQUEOS", f"C{est_row}"))
    print(f"  conversión a venta en 0: inversión para la meta = {meta_val!r}, Chequeos {e4}")
    if not str(meta_val).startswith("inalcanzable") or str(e4) != "REVISAR":
        fallas.append(f"meta inalcanzable mal señalada: {meta_val!r} / Chequeos {e4}")
    # resumen
    if fallas:
        print("  FALLAS:")
        for f in fallas[:25]:
            print("   -", f)
    else:
        print("  TODO OK")
    return not fallas


def verificar_simulador(path, v, wb):
    """Con los supuestos sin tocar, cada fila del Simulador tiene que dar lo mismo que el Modelo; y al cambiar un supuesto, cambiar."""
    if "Simulador" not in wb.sheetnames:
        return []
    ws = wb["Simulador"]; fallas = []; comparadas = 0; fila_in = None; fila_profit = None
    for r in range(5, ws.max_row + 1):
        e, f = ws.cell(r, 5).value, ws.cell(r, 6).value
        if isinstance(e, str) and e.startswith("=") and isinstance(f, str) and f.startswith("="):
            ve, vf = v.get(("SIMULADOR", f"E{r}")), v.get(("SIMULADOR", f"F{r}")); comparadas += 1
            iguales = (abs(float(ve) - float(vf)) < 1e-6) if isinstance(ve, (int, float)) and isinstance(vf, (int, float)) else (str(ve) == str(vf))
            if not iguales: fallas.append(f"Simulador fila {r} ({ws.cell(r, 2).value}): {ve} ≠ Modelo {vf}")
            if ws.cell(r, 2).value == "Ganancia del período": fila_profit = r
        c = ws.cell(r, 3).value
        if fila_in is None and isinstance(c, str) and c.startswith("='Supuestos'"):
            fila_in = r   # la primera celda amarilla del Simulador (la inversión o el CPM)
    cambia = None
    if fila_in and fila_profit:
        import tempfile
        wb2 = openpyxl.load_workbook(path); ws2 = wb2["Simulador"]
        base = v.get(("SIMULADOR", f"D{fila_in}")); ws2.cell(fila_in, 3).value = float(base) * 1.5
        tmp = os.path.join(tempfile.gettempdir(), "sim_test.xlsx"); wb2.save(tmp); v2 = calcular(tmp)
        antes, despues = v.get(("SIMULADOR", f"E{fila_profit}")), v2.get(("SIMULADOR", f"E{fila_profit}"))
        cambia = abs(float(antes) - float(despues)) > 1e-6
        if not cambia: fallas.append(f"el Simulador no reacciona al cambiar el primer supuesto (ganancia {antes} → {despues})")
    # un porcentaje imposible en el Simulador tiene que poner Chequeos en REVISAR
    detecta = None; fila_pct = None
    for r in range(5, ws.max_row + 1):
        c = ws.cell(r, 3)
        if isinstance(c.value, str) and c.value.startswith("='Supuestos'") and "%" in (c.number_format or ""):
            fila_pct = r; break
    if fila_pct:
        import tempfile
        wb3 = openpyxl.load_workbook(path); wb3["Simulador"].cell(fila_pct, 3).value = 1.5
        tmp3 = os.path.join(tempfile.gettempdir(), "sim_invalido.xlsx"); wb3.save(tmp3); v3 = calcular(tmp3)
        chq = wb3["Chequeos"]; fila_general = next((rr for rr in range(5, chq.max_row + 1) if chq.cell(rr, 2).value == "Estado general del modelo"), None)
        estado = v3.get(("CHEQUEOS", f"C{fila_general}")) if fila_general else None
        detecta = (str(estado) == "REVISAR")
        if not detecta: fallas.append(f"Chequeos dice {estado} con un porcentaje de 150 % en el Simulador (fila {fila_pct})")
    print(f"  Simulador: {comparadas} filas comparadas con Modelo, {len(fallas)} diferencias; reacciona al cambiar el primer supuesto: {cambia}; Chequeos detecta un valor inválido: {detecta}")
    return fallas


def verificar_simulacion(sim_path, modelo_path):
    """El Excel de simulación, con los valores del ejemplo, tiene que dar lo mismo que la hoja Modelo del modelo; reaccionar
    al cambiar un supuesto; y marcar REVISAR si se escribe un porcentaje imposible."""
    import tempfile
    print(f"\n=== {os.path.basename(sim_path)}  (contra {os.path.basename(modelo_path)})")
    wb_s = openpyxl.load_workbook(sim_path); ws = wb_s["Simulación"]; vs = calcular(sim_path)
    wb_m = openpyxl.load_workbook(modelo_path); wm = wb_m["Modelo"]; vm = calcular(modelo_path)
    fallas = []; n_f = 0
    for hoja in wb_s.worksheets:
        for row in hoja.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    n_f += 1; v = vs.get((hoja.title.upper(), c.coordinate))
                    if v is None or es_error(v): fallas.append(f"{hoja.title}!{c.coordinate} → {v}")
    modelo = {}
    for r in range(5, wm.max_row + 1):
        lab, f = wm.cell(r, 2).value, wm.cell(r, 3).value
        if lab and isinstance(f, str) and f.startswith("="): modelo[lab] = vm.get(("MODELO", f"C{r}"))
    comparadas = 0; fila_in = fila_pct = fila_profit = None
    for r in range(6, ws.max_row + 1):
        lab, c = ws.cell(r, 2).value, ws.cell(r, 3)
        if isinstance(c.value, str) and c.value.startswith("=") and lab in modelo:
            a, b = vs.get(("SIMULACIÓN", f"C{r}")), modelo[lab]; comparadas += 1
            iguales = (abs(float(a) - float(b)) < 1e-6) if isinstance(a, (int, float)) and isinstance(b, (int, float)) else (str(a) == str(b))
            if not iguales: fallas.append(f"fila {r} ({lab}): simulación {a} ≠ Modelo {b}")
            if lab == "Ganancia del período": fila_profit = r
        elif isinstance(c.value, (int, float)):
            if fila_in is None: fila_in = r
            if fila_pct is None and "%" in (c.number_format or ""): fila_pct = r
    reacciona = detecta = None
    if fila_in and fila_profit:
        wb2 = openpyxl.load_workbook(sim_path); wb2["Simulación"].cell(fila_in, 3).value = float(wb2["Simulación"].cell(fila_in, 3).value) * 1.5
        tmp = os.path.join(tempfile.gettempdir(), "simulacion_test.xlsx"); wb2.save(tmp); v2 = calcular(tmp)
        reacciona = abs(float(vs.get(("SIMULACIÓN", f"C{fila_profit}"))) - float(v2.get(("SIMULACIÓN", f"C{fila_profit}")))) > 1e-6
        if not reacciona: fallas.append("la simulación no reacciona al cambiar el primer supuesto")
    if fila_pct:
        wb3 = openpyxl.load_workbook(sim_path); wb3["Simulación"].cell(fila_pct, 3).value = 1.5
        tmp3 = os.path.join(tempfile.gettempdir(), "simulacion_invalida.xlsx"); wb3.save(tmp3); v3 = calcular(tmp3)
        detecta = str(v3.get(("SIMULACIÓN", "C4"))) == "REVISAR"
        if not detecta: fallas.append(f"'Estado de tus números' no marca REVISAR con 150 % en la fila {fila_pct}")
    # una cobranza que suma más de 100 % (infla el cash) tiene que marcar REVISAR aunque cada celda esté en rango
    detecta2 = None; fila_m1 = next((r for r in range(6, ws.max_row + 1) if str(ws.cell(r, 2).value or "").startswith("% del facturado que entra en el mes 1")), None)
    if fila_m1:
        wb4 = openpyxl.load_workbook(sim_path); wb4["Simulación"].cell(fila_m1, 3).value = 0.95
        tmp4 = os.path.join(tempfile.gettempdir(), "simulacion_cobranza.xlsx"); wb4.save(tmp4); v4 = calcular(tmp4)
        detecta2 = str(v4.get(("SIMULACIÓN", "C4"))) == "REVISAR"
        if not detecta2: fallas.append("'Estado de tus números' no marca REVISAR con una cobranza que suma más de 100 %")
    n_chq = sum(1 for r in range(6, ws.max_row + 1) if isinstance(ws.cell(r, 3).value, str) and ws.cell(r, 3).value.startswith('=IF(') and ws.cell(r, 3).value.endswith('"OK","REVISAR")'))
    print(f"  fórmulas: {n_f} · filas comparadas con Modelo: {comparadas} · diferencias: {sum(1 for f in fallas if '≠' in f)} · reacciona: {reacciona} · detecta % inválido: {detecta} · detecta cobranza > 100 %: {detecta2} · chequeos de consistencia: {n_chq}")
    for f in fallas[:8]: print("   -", f)
    return not fallas


def main():
    carpeta = sys.argv[1] if len(sys.argv) > 1 else DEFAULT_IN
    ok = True
    for p in sorted(glob.glob(os.path.join(carpeta, "*.xlsx"))):
        ok = verificar(p) and ok
    for sp in sorted(glob.glob(os.path.join(carpeta, "simuladores", "*.xlsx"))):
        pref = os.path.basename(sp)[:3]
        modelos = [m for m in glob.glob(os.path.join(carpeta, "*.xlsx")) if os.path.basename(m).startswith(pref)]
        if modelos: ok = verificar_simulacion(sp, modelos[0]) and ok
    print("\nRESULTADO GENERAL:", "OK" if ok else "CON FALLAS")
    sys.exit(0 if ok else 1)


if __name__ == "__main__":
    main()
